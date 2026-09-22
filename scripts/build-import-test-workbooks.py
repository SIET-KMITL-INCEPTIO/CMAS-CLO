# -*- coding: utf-8 -*-
"""ไฟล์ทดสอบสำหรับตัวนำเข้าของ docs/pages/index-q.html

   ไฟล์เหล่านี้เลียนเทมเพลตที่ระบบสร้างจริง — หัวเอกสาร 12 แถว หัวตารางแถว 13
   ข้อมูลเริ่มแถว 14 พร้อมตรึงหัว data validation และการล็อกช่อง — เพราะเส้นทางที่
   ต้องทดสอบคือ "ดาวน์โหลดเทมเพลต → กรอกใน Excel → นำเข้ากลับ" ไม่ใช่ไฟล์แบนที่
   ไม่มีใครได้รับจริง

     01-scores-valid        ทุกแถวผ่าน · เว้นบางช่องไว้เพื่อพิสูจน์ CR-03.1
     02-scores-errors       หนึ่งแถวต่อหนึ่งกติกาการปฏิเสธ + หนึ่งแถวที่ถูกต้อง
     03-roster              รายชื่อใหม่ + เคสซ้ำและเคสข้อมูลขาด
     04-legacy-flat         ไฟล์แบนไม่มีหัวเอกสาร — พิสูจน์ว่าไฟล์รุ่นเก่ายังนำเข้าได้

   รายชื่อและรหัสกิจกรรมด้านล่างสะท้อนรายวิชา c1 (90641003 Sec 01) ใน
   docs/pages/index-q.html · ถ้า fixture นั้นเปลี่ยน ให้รันสคริปต์นี้ใหม่ด้วยค่าใหม่ —
   ข้อความในวงเล็บเหลี่ยมท้ายหัวคอลัมน์คือสิ่งที่ผูกคอลัมน์กับกิจกรรม รหัสเก่าทำให้
   ไฟล์อ่านไม่ออกทั้งไฟล์

     python scripts/build-import-test-workbooks.py
"""
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parents[1] / 'docs' / 'excel' / 'import-test'

ORG = {'institute': 'มหาวิทยาลัย', 'faculty': 'คณะวิศวกรรมศาสตร์'}

COURSE = {
    'code': '90641003', 'name': 'การวิเคราะห์และออกแบบระบบสารสนเทศ',
    'section': '01', 'semester': 1, 'year': 2568,
}

STAFF = 'ผศ.ดร. สมชาย ใจดี (ผู้ประสานงานรายวิชา) · อ. กัญญา ศรีสุข (ผู้สอนร่วม)'
AUTHOR = 'ผศ.ดร. สมชาย ใจดี (ผู้ประสานงานรายวิชา)'

ACTS = [
    ('a11', 'แบบทดสอบย่อย ครั้งที่ 1', 20, 10),
    ('a12', 'สอบกลางภาค', 100, 25),
    ('a13', 'ปฏิบัติการ UML', 50, 20),
    ('a14', 'โครงงานกลุ่ม', 100, 35),
    ('a15', 'การมีส่วนร่วมในชั้นเรียน', 10, 10),
]

ROSTER = [
    ('6604101001', 'กชกร วงศ์อารีย์'),   ('6604101002', 'ธนกฤต บุญมาก'),
    ('6604101003', 'ปิยะดา แสนสุข'),     ('6604101004', 'ณัฐวุฒิ พงษ์ไพร'),
    ('6604101005', 'อารยา ทองแท้'),      ('6604101006', 'สหรัฐ เกษมสุข'),
    ('6604101007', 'มนัสนันท์ ใจงาม'),   ('6604101008', 'ภาคิน รุ่งเรือง'),
    ('6604101009', 'ชนากานต์ ดีพร้อม'),  ('6604101010', 'วรินทร ศรีสมบัติ'),
    ('6604101011', 'อนุชา จันทร์เพ็ญ'),  ('6604101012', 'เบญจวรรณ ปานทอง'),
]

# ── สไตล์ ให้ตรงกับ STYLES_XML ในต้นแบบ ──────────────────────────────
TITLE_FONT = Font(name='Tahoma', size=15, bold=True, color='FF1E233C')
LABEL_FONT = Font(name='Tahoma', size=10, color='FF6B7192')
HEAD_FONT = Font(name='Tahoma', size=11, bold=True, color='FFFFFFFF')
NOTE_FONT = Font(name='Tahoma', size=9, color='FF878CB4')
HEAD_FILL = PatternFill('solid', fgColor='FF1E233C')
IDENT_FILL = PatternFill('solid', fgColor='FFF1F2F8')
THIN = Side(style='thin', color='FFD8DAE8')
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
LOCKED = Protection(locked=True)
OPEN = Protection(locked=False)

TH_MONTH = ['มกราคม', 'กุมภาพันธ์', 'มีนาคม', 'เมษายน', 'พฤษภาคม', 'มิถุนายน',
            'กรกฎาคม', 'สิงหาคม', 'กันยายน', 'ตุลาคม', 'พฤศจิกายน', 'ธันวาคม']


def thai_date(d):
    return '%d %s %d เวลา %02d:%02d น.' % (d.day, TH_MONTH[d.month - 1], d.year + 543, d.hour, d.minute)


def stamp(d):
    return '%d%02d%02d-%02d%02d' % (d.year + 543, d.month, d.day, d.hour, d.minute)


def doc_id(kind, d):
    return '-'.join(['CMAS', kind.upper(), COURSE['code'], COURSE['section'],
                     '%d-%d' % (COURSE['semester'], COURSE['year']), stamp(d)])


def col_head(a):
    return '%s (เต็ม %d) [%s]' % (a[1], a[2], a[0])


def write_header(ws, title, kind, ncols, now):
    """หัวเอกสาร 12 แถว · หัวตารางจะไปอยู่แถว 13 เหมือนที่ระบบสร้าง"""
    rows = [
        (title, None), (ORG['institute'] + ' · ' + ORG['faculty'], None), (None, None),
        ('รายวิชา', COURSE['code'] + ' ' + COURSE['name']),
        ('Sec', COURSE['section']),
        ('ภาค/ปีการศึกษา', '%d/%d' % (COURSE['semester'], COURSE['year'])),
        ('ผู้สอน', STAFF),
        ('วันที่สร้าง', thai_date(now)),
        ('ผู้สร้างเอกสาร', AUTHOR),
        ('รหัสเอกสาร', doc_id(kind, now)),
        (None, None), (None, None),
    ]
    for r, (a, b) in enumerate(rows, start=1):
        if a is None:
            continue
        ws.cell(row=r, column=1, value=a)
        if b is None:                                   # แถวชื่อเอกสาร
            ws.cell(row=r, column=1).font = TITLE_FONT if r == 1 else LABEL_FONT
            if ncols > 1:
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        else:
            ws.cell(row=r, column=1).font = LABEL_FONT
            ws.cell(row=r, column=2, value=b)
            if ncols > 2:
                ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=ncols)
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 15
    return 13                                           # แถวหัวตาราง


def write_table(ws, hdr_row, head, rows, score_from=2):
    for c, v in enumerate(head, start=1):
        cell = ws.cell(row=hdr_row, column=c, value=v)
        cell.font, cell.fill, cell.border = HEAD_FONT, HEAD_FILL, GRID
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for r, row in enumerate(rows, start=hdr_row + 1):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=(None if v == '' else v))
            cell.border = GRID
            if c <= score_from:                          # คอลัมน์ระบุตัวตน
                cell.fill, cell.protection = IDENT_FILL, LOCKED
            else:                                        # ช่องคะแนน
                cell.protection = OPEN
                cell.number_format = '0.##'
                cell.alignment = Alignment(horizontal='right')
    ws.freeze_panes = '%s%d' % (col_letter(score_from + 1), hdr_row + 1)


def add_validations(ws, hdr_row, nrows):
    for i, a in enumerate(ACTS):
        col = col_letter(3 + i)
        dv = DataValidation(
            type='decimal', operator='between', formula1='0', formula2=str(a[2]),
            allow_blank=True, showInputMessage=True, showErrorMessage=True, errorStyle='stop',
            errorTitle='คะแนนไม่ถูกต้อง',
            error='กรอกได้เฉพาะตัวเลข 0 ถึง %d · เว้นว่างได้ถ้ายังไม่ประเมิน' % a[2],
            promptTitle=a[1],
            prompt='เต็ม %d คะแนน · เว้นว่าง = ยังไม่ประเมิน' % a[2])
        dv.add('%s%d:%s%d' % (col, hdr_row + 1, col, hdr_row + nrows))
        ws.add_data_validation(dv)


def col_letter(i):
    """เลขคอลัมน์ (1-based) เป็นตัวอักษร · ใช้ ws.cell().column_letter ไม่ได้เพราะ
       เซลล์ในแถวที่ถูก merge เป็น MergedCell ซึ่งไม่มี attribute นั้น"""
    name = ''
    while i > 0:
        i, m = divmod(i - 1, 26)
        name = chr(65 + m) + name
    return name


def guide_sheet(wb, lines, widths=(26, 66)):
    ws = wb.create_sheet('คำแนะนำ')
    for ln in lines:
        ws.append(list(ln))
    ws['A1'].font = Font(name='Tahoma', bold=True, size=13)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[col_letter(i)].width = w
    return ws


def finish(ws, widths, protect=True):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[col_letter(i)].width = w
    if protect:
        # ไม่ตั้งรหัสผ่าน — เป็นที่กั้นไม่ให้พิมพ์ผิดช่อง ไม่ใช่ความปลอดภัย
        # (openpyxl โยน TypeError ถ้าตั้ง password = None ตรง ๆ)
        ws.protection.sheet = True
        ws.protection.formatCells = False
        ws.protection.insertRows = False
        ws.protection.deleteRows = False
        ws.protection.selectLockedCells = False
        ws.protection.selectUnlockedCells = False


def props(wb, title, kind, now):
    wb.properties.title = title
    wb.properties.subject = COURSE['name']
    wb.properties.creator = 'ผศ.ดร. สมชาย ใจดี'
    wb.properties.category = ORG['institute'] + ' · ' + ORG['faculty']
    wb.properties.keywords = doc_id(kind, now)


# ══ 01 · ทุกแถวผ่าน ═══════════════════════════════════════════════════
def scores_valid(now):
    wb = Workbook()
    ws = wb.active
    ws.title = 'คะแนน'
    head = ['รหัสนักศึกษา', 'ชื่อ-นามสกุล'] + [col_head(a) for a in ACTS]
    pattern = [
        [18, 82, 46, 91, 9], [15, 74, 41, 86, 8], [12, 63, 35, 78, 7],
        [19, 88, 48, 95, 10], [9, 55, 28, 64, 6], [16, 71, 39, 83, 8],
        [11, 60, 33, '', 7], [17, 79, 44, '', 9], [8, 48, 25, '', 5],
        [14, 68, 37, 80, 7], [20, 92, 50, 97, 10], [13, 66, 36, 76, 8],
    ]
    rows = [[c, n] + pattern[i] for i, (c, n) in enumerate(ROSTER)]
    hdr = write_header(ws, 'แบบฟอร์มกรอกคะแนน (ไฟล์ทดสอบ 01)', 'template', len(head), now)
    write_table(ws, hdr, head, rows)
    add_validations(ws, hdr, len(rows))
    finish(ws, [15, 30] + [20] * len(ACTS))
    guide_sheet(wb, [
        ['ไฟล์ทดสอบ 01 — คะแนนถูกต้องทั้งหมด'], [],
        ['รูปแบบไฟล์', 'ตรงกับเทมเพลตที่ระบบสร้าง — หัวเอกสารแถว 1–12 หัวตารางแถว 13 ข้อมูลเริ่มแถว 14'],
        ['คาดหวังเมื่อนำเข้า', 'ผ่าน 57 ช่อง · ไม่ผ่าน 0 รายการ'],
        ['จุดที่ต้องสังเกต', 'แถว 20–22 เว้นช่อง “โครงงานกลุ่ม” ไว้'],
        ['', 'ช่องว่าง = ไม่มีข้อมูลมาด้วย ระบบจึงไม่แตะช่องนั้น — ไม่เขียนทับคะแนนเดิม'],
        ['', 'และไม่บันทึกเป็น 0 ถ้าเดิมยังไม่เคยมีคะแนน (CR-03.1)'],
        ['', 'ทดสอบผลข้อนี้ได้ชัดที่สุดกับ Sec ที่ยังไม่มีคะแนนเลย'],
        ['', 'คะแนนรวมคิดจากน้ำหนักที่ประเมินแล้วเท่านั้น (CR-05)'], [],
        ['ลองใน Excel', 'พิมพ์ 999 ในช่องคะแนน — Excel จะปฏิเสธเองก่อนถึงระบบ'],
        ['', 'ลองแก้ช่องรหัสหรือชื่อ — ถูกล็อกไว้ด้วย sheet protection'],
    ])
    props(wb, 'ไฟล์ทดสอบ 01 คะแนนถูกต้อง', 'template', now)
    return wb, '01-scores-valid.xlsx'


# ══ 02 · ครบทุกกติกาการปฏิเสธ ═════════════════════════════════════════
def scores_errors(now):
    wb = Workbook()
    ws = wb.active
    ws.title = 'คะแนน'
    head = ['รหัสนักศึกษา', 'ชื่อ-นามสกุล'] + [col_head(a) for a in ACTS]
    blank = ['', '', '', '']
    rows = [
        [ROSTER[0][0], ROSTER[0][1], -3] + blank,                    # แถว 14 ติดลบ
        [ROSTER[1][0], ROSTER[1][1], 25] + blank,                    # แถว 15 เกินเต็ม 20
        [ROSTER[2][0], ROSTER[2][1], 'ไม่มา'] + blank,               # แถว 16 ไม่ใช่ตัวเลข
        ['6699999999', 'ผี ไม่มีใน Sec', 10] + blank,            # แถว 17 ไม่มีรหัสนี้
        [ROSTER[0][0], 'กชกร (แถวซ้ำ)', 12] + blank,                 # แถว 18 รหัสซ้ำ
        [ROSTER[3][0], ROSTER[3][1], 17, 80, 44, 88, 9],             # แถว 19 ถูกต้อง
        ['', '', '', '', '', '', ''],                                # แถว 20 ว่างทั้งแถว
    ]
    hdr = write_header(ws, 'ไฟล์ทดสอบ 02 — ข้อมูลผิดครบทุกแบบ', 'scores', len(head), now)
    write_table(ws, hdr, head, rows)
    # ตั้งใจไม่ใส่ data validation ในไฟล์นี้ — ค่าที่ผิดต้องอยู่ในไฟล์ให้ได้
    # จึงจะทดสอบด่านตรวจฝั่งระบบ ซึ่งเป็นด่านที่ต้องเชื่อถือได้จริง
    finish(ws, [15, 30] + [20] * len(ACTS), protect=False)
    guide_sheet(wb, [
        ['ไฟล์ทดสอบ 02 — ครบทุกกติกาการปฏิเสธ'], [],
        ['รูปแบบไฟล์', 'หัวเอกสารแถว 1–12 หัวตารางแถว 13 ข้อมูลเริ่มแถว 14'],
        ['คาดหวังเมื่อนำเข้า', 'ผ่าน 5 ช่อง (แถว 19) · ไม่ผ่าน 5 รายการ'],
        ['', 'เลขแถวในรายงานต้องตรงกับเลขแถวที่เห็นในไฟล์นี้'], [],
        ['แถว 14', 'คะแนนติดลบ (-3)'],
        ['แถว 15', 'เกินคะแนนเต็ม (25 จากเต็ม 20)'],
        ['แถว 16', 'ไม่ใช่ตัวเลข ("ไม่มา")'],
        ['แถว 17', 'ไม่มีรหัสนี้ใน Sec'],
        ['แถว 18', 'รหัสซ้ำกับแถว 14'],
        ['แถว 19', 'ถูกต้อง — ต้องถูกบันทึกแม้แถวอื่นถูกปฏิเสธ'],
        ['แถว 20', 'ว่างทั้งแถว — ต้องถูกข้ามโดยไม่นับเป็นข้อผิดพลาด'], [],
        ['ทำไมไฟล์นี้ไม่มี data validation', 'ค่าที่ผิดต้องอยู่ในไฟล์ให้ได้ จึงจะทดสอบด่านตรวจฝั่งระบบ'],
        ['', 'ซึ่งเป็นด่านที่ต้องเชื่อถือได้จริง · validation ใน Excel เป็นเพียงด่านแรก'], [],
        ['หลังนำเข้า', 'ดูรายงานได้ที่ นักศึกษาและคะแนน › ประวัติการนำเข้า'],
    ])
    props(wb, 'ไฟล์ทดสอบ 02 ข้อมูลผิด', 'scores', now)
    return wb, '02-scores-errors.xlsx'


# ══ 03 · รายชื่อนักศึกษา ══════════════════════════════════════════════
def roster(now):
    wb = Workbook()
    ws = wb.active
    ws.title = 'รายชื่อ'
    head = ['รหัสนักศึกษา', 'ชื่อ-นามสกุล']
    rows = [
        ['6604101021', 'ศิรินทิพย์ กล้าหาญ'],          # แถว 14
        ['6604101022', 'พงศกร อินทรีย์'],              # แถว 15
        ['6604101023', 'ญาณิศา เพชรรัตน์'],            # แถว 16
        ['6604101021', 'ศิรินทิพย์ (แถวซ้ำในไฟล์)'],    # แถว 17
        [ROSTER[0][0], 'กชกร (มีในระบบแล้ว)'],          # แถว 18
        ['', 'ไม่ได้กรอกรหัส'],                        # แถว 19
        ['6604101024', ''],                            # แถว 20
    ]
    hdr = write_header(ws, 'ไฟล์ทดสอบ 03 — นำเข้ารายชื่อนักศึกษา', 'roster', len(head), now)
    write_table(ws, hdr, head, rows, score_from=0)     # ทั้งสองคอลัมน์กรอกได้
    finish(ws, [18, 38], protect=False)
    guide_sheet(wb, [
        ['ไฟล์ทดสอบ 03 — นำเข้ารายชื่อนักศึกษา'], [],
        ['รูปแบบไฟล์', 'หัวเอกสารแถว 1–12 หัวตารางแถว 13 ข้อมูลเริ่มแถว 14'],
        ['คาดหวังเมื่อนำเข้า', 'เพิ่ม 3 คน · ไม่ผ่าน 4 รายการ'], [],
        ['แถว 17', 'รหัสซ้ำกับแถว 14 ในไฟล์เดียวกัน'],
        ['แถว 18', 'มีรหัสนี้ใน Sec อยู่แล้ว — ต้องไม่ทับชื่อเดิม'],
        ['แถว 19', 'ไม่มีรหัสนักศึกษา'],
        ['แถว 20', 'ไม่มีชื่อ'], [],
        ['หลังนำเข้า', 'สามคนใหม่ต้องขึ้น “ยังประเมินไม่ครบ” ทุกกิจกรรม'],
        ['', 'และต้องปรากฏในรายการ “นักศึกษาที่ต้องติดตาม” หน้าภาพรวม'],
    ])
    props(wb, 'ไฟล์ทดสอบ 03 รายชื่อนักศึกษา', 'roster', now)
    return wb, '03-roster.xlsx'


# ══ 04 · ไฟล์แบนรุ่นเก่า ═══════════════════════════════════════════════
def legacy_flat(now):
    """ไม่มีหัวเอกสาร หัวตารางอยู่แถว 1 — ไฟล์แบบที่ระบบเคยสร้างก่อนรอบนี้
       และแบบที่คนสร้างเองจาก Google Sheets · ตัวนำเข้าหาหัวตารางจากเนื้อหา
       จึงต้องอ่านได้เท่ากัน ข้อนี้ถูกอ้างว่ารองรับ จึงต้องมีไฟล์ทดสอบยืนยัน"""
    wb = Workbook()
    ws = wb.active
    ws.title = 'คะแนน'
    head = ['รหัสนักศึกษา', 'ชื่อ-นามสกุล'] + [col_head(a) for a in ACTS]
    ws.append(head)
    ws.append([ROSTER[0][0], ROSTER[0][1], 11, 61, 31, 71, 6])      # แถว 2
    ws.append([ROSTER[1][0], ROSTER[1][1], 12, 62, 32, 72, 7])      # แถว 3
    ws.append([ROSTER[2][0], ROSTER[2][1], -1, '', '', '', ''])     # แถว 4 ติดลบ
    finish(ws, [15, 30] + [20] * len(ACTS), protect=False)
    guide_sheet(wb, [
        ['ไฟล์ทดสอบ 04 — ไฟล์แบนไม่มีหัวเอกสาร'], [],
        ['ทำไมต้องมีไฟล์นี้', 'ระบบอ้างว่ารองรับไฟล์รุ่นเก่าและไฟล์ที่คนทำเอง คำอ้างนั้นต้องมีไฟล์ทดสอบยืนยัน'],
        ['รูปแบบไฟล์', 'หัวตารางอยู่แถว 1 ข้อมูลเริ่มแถว 2 ไม่มีหัวเอกสาร ไม่มีการจัดรูปแบบ'],
        ['คาดหวังเมื่อนำเข้า', 'ผ่าน 10 ช่อง · ไม่ผ่าน 1 รายการ (แถว 4 คะแนนติดลบ)'],
    ])
    props(wb, 'ไฟล์ทดสอบ 04 ไฟล์แบน', 'scores', now)
    return wb, '04-legacy-flat.xlsx'


def main():
    now = datetime.now()
    OUT.mkdir(parents=True, exist_ok=True)
    for build in (scores_valid, scores_errors, roster, legacy_flat):
        wb, name = build(now)
        wb.save(OUT / name)
        print('wrote', (OUT / name).relative_to(OUT.parents[2]))


if __name__ == '__main__':
    main()
