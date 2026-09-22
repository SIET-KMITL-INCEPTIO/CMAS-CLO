"""
Build docs/excel/CMAS-ER-Tables.xlsx — one sheet per table, every column, with
PK / FK / UNIQUE highlighted.

Layout per sheet keeps the three columns the advisor already reads and adds two:

    คีย์ | หัวข้อ | ค่าที่กรอกในระบบ | เก็บที่คอลัมน์ | ชนิดข้อมูล

SELF-CHECKING: at build time every column below is compared against
docs/reference/db/mysql/cmas_app_mysql_v4.sql — the same file MySQL Workbench
reverse-engineers into the ER diagram. Column names, order, types, PK, FK and
UNIQUE must all match or the build fails loudly. So the workbook cannot drift
from the diagram: if the schema changes and this file is not updated, nothing is
produced at all.

Run: python scripts/build-er-tables-workbook.py
Output: docs/excel/CMAS-ER-Tables.xlsx
"""

import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ROOT = Path(__file__).resolve().parent.parent
SQL = ROOT / "docs" / "reference" / "db" / "mysql" / "cmas_app_mysql_v4.sql"
OUT = ROOT / "docs" / "excel" / "CMAS-ER-Tables.xlsx"

# ---------------------------------------------------------------------------
# Palette — same tokens as the other workbooks in this folder
# ---------------------------------------------------------------------------
C_PRIMARY = "19151B"
C_UI = "1A56C4"
C_ON = "1A1C1D"
C_MUTED = "4A454A"
C_OUTLINE = "CCC4CA"
C_CALC = "137333"
C_TQF = "8A5A00"
C_DERIV = "44405A"
TX_GAP = "C5221F"

BG_UI = "EAF1FE"
BG_CALC = "E6F4EA"
BG_TQF = "FEF7E0"
BG_ROW = "F3F3F5"
BG_GAP = "F8CBAD"
BG_DERIV = "EDEBF2"

thin = Side(style="thin", color=C_OUTLINE)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

F_TITLE = Font(name="Calibri", size=16, bold=True, color=C_UI)
F_SUB = Font(name="Calibri", size=10.5, italic=True, color=C_MUTED)
F_H2 = Font(name="Calibri", size=12, bold=True, color=C_PRIMARY)
F_LABEL = Font(name="Calibri", size=10, bold=True, color=C_ON)
F_BODY = Font(name="Calibri", size=10, color=C_ON)
F_MUTED = Font(name="Calibri", size=9.5, color=C_MUTED)
F_HEAD = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
F_MONO = Font(name="Consolas", size=9.5, color=C_MUTED)
F_TYPE = Font(name="Consolas", size=9, color=C_MUTED)
F_GAP = Font(name="Calibri", size=9.5, bold=True, color=TX_GAP)
F_DERIVF = Font(name="Consolas", size=9.5, italic=True, color=C_DERIV)

FILL_UI = PatternFill("solid", fgColor=C_UI)
FILL_INPUT = PatternFill("solid", fgColor=BG_UI)
FILL_AUTO = PatternFill("solid", fgColor=BG_CALC)
FILL_ROW = PatternFill("solid", fgColor=BG_ROW)
FILL_WHITE = PatternFill("solid", fgColor="FFFFFF")
FILL_GAP = PatternFill("solid", fgColor=BG_GAP)
FILL_DERIV = PatternFill("solid", fgColor=BG_DERIV)

# คีย์ badge styling: (font colour, fill colour)
KEY_STYLE = {
    "PK": (C_CALC, BG_CALC),
    "FK": (C_UI, BG_UI),
    "UK": (C_TQF, BG_TQF),
    "FK · UK": (C_UI, BG_UI),
    "": (C_MUTED, "FFFFFF"),
}

IN, AUTO, DERIV, GAP = "in", "auto", "deriv", "gap"

COURSE_LINE = "03376120 ระบบฐานข้อมูล (ภาค 1/2568 หมู่ 01)"

# ---------------------------------------------------------------------------
# THE SCHEMA, as the workbook presents it.
# (column, type, key, หัวข้อ, ค่าตัวอย่าง, kind, ref)  — ref = target of an FK
# Rows whose `column` is None are NOT columns: they are values the screens show
# but nothing stores (DERIV), or things มคอ.3 asks for that v1 has no room for
# (GAP). They are excluded from the comparison against the SQL.
# ---------------------------------------------------------------------------
TABLES = [
    ("User", "ผู้ใช้ระบบ — ADMIN หรือ INSTRUCTOR (ไม่มี role นักศึกษา)", "1 แถว = 1 บัญชี", [
        ("id", "VARCHAR(30)", "PK", "รหัสผู้ใช้ในระบบ", "usr-02", AUTO, None),
        ("email", "VARCHAR(255)", "UK", "อีเมล (ใช้เข้าสู่ระบบ)", "wichai.k@kmitl.ac.th", IN, None),
        ("name", "VARCHAR(255)", "", "ชื่อ-นามสกุล", "ผศ.วิชัย คงเจริญ", IN, None),
        ("passwordHash", "VARCHAR(255)", "", "รหัสผ่าน",
         "•••••••• (เก็บเป็นค่าแฮช argon2 ไม่เก็บรหัสผ่านจริง)", IN, None),
        ("role", "ENUM(2)", "", "สิทธิ์การใช้งาน", "INSTRUCTOR (อีกค่าคือ ADMIN)", IN, None),
        ("isActive", "TINYINT(1)", "", "สถานะใช้งาน", "ใช้งานอยู่ (ปิดบัญชีแทนการลบเสมอ)", IN, None),
        ("createdAt", "DATETIME(3)", "", "วันที่สร้างบัญชี", "2025-06-01 09:00", AUTO, None),
        ("updatedAt", "DATETIME(3)", "", "วันที่แก้ไขล่าสุด", "2025-06-01 09:00", AUTO, None),
    ]),

    ("Course", "รายวิชา — รากของลำดับชั้นข้อมูลทั้งหมด ไม่มีอะไรอยู่เหนือกว่านี้",
     "1 แถว = 1 หมู่เรียน", [
        ("id", "VARCHAR(30)", "PK", "รหัสรายวิชาในระบบ", "crs-01", AUTO, None),
        ("code", "VARCHAR(50)", "UK", "รหัสวิชา", "03376120", IN, None),
        ("name", "VARCHAR(255)", "", "ชื่อรายวิชา (ไทย)", "ระบบฐานข้อมูล", IN, None),
        ("nameEn", "VARCHAR(255)", "", "ชื่อรายวิชา (อังกฤษ)", "DATABASE SYSTEM", IN, None),
        ("semester", "INT", "UK", "ภาคการศึกษา", "1", IN, None),
        ("year", "INT", "UK", "ปีการศึกษา (พ.ศ.)", "2568", IN, None),
        ("section", "VARCHAR(10)", "UK", "กลุ่มเรียน (หมู่เรียน)", "01", IN, None),
        ("credits", "DECIMAL(3,1)", "", "จำนวนหน่วยกิต", "3", IN, None),
        ("lectureHours", "DECIMAL(4,1)", "", "ชั่วโมงบรรยาย / สัปดาห์", "2", IN, None),
        ("practiceHours", "DECIMAL(4,1)", "", "ชั่วโมงปฏิบัติ / สัปดาห์", "2", IN, None),
        ("selfStudyHours", "DECIMAL(4,1)", "", "ชั่วโมงศึกษาด้วยตนเอง / สัปดาห์", "5", IN, None),
        ("gradeScale", "ENUM(2)", "", "รูปแบบผลการเรียน",
         "LETTER (A–F) · อีกค่าคือ PASS_FAIL (S/U)", IN, None),
        ("gradeMethod", "ENUM(2)", "", "วิธีตัดเกรด",
         "อิงเกณฑ์ (CRITERION_REFERENCED) · อีกค่าคือ อิงกลุ่ม (NORM_REFERENCED)", IN, None),
        ("passCriteria", "DOUBLE", "", "เกณฑ์คะแนนรวมที่ถือว่าผ่านรายวิชา (%)", "60", IN, None),
        ("classTarget", "DOUBLE", "", "เป้าหมายสัดส่วนผู้ผ่านต่อ CLO (%)", "70", IN, None),
        ("createdAt", "DATETIME(3)", "", "วันที่สร้าง", "2025-06-01 09:30", AUTO, None),
        ("updatedAt", "DATETIME(3)", "", "วันที่แก้ไขล่าสุด", "2025-06-01 09:30", AUTO, None),
        (None, "", "", "หน่วยของเส้นแบ่งเกรด", "ร้อยละ เมื่อตัดอิงเกณฑ์ · T-score เมื่อตัดอิงกลุ่ม",
         DERIV, "ตามมาจาก gradeMethod ไม่ต้องเก็บ"),
        (None, "", "", "หลักสูตรและประเภทของรายวิชา",
         "หมวดวิชาเฉพาะ · กลุ่มวิชาชีพเฉพาะสาขาวิชา (บังคับเรียน)", GAP,
         "ไม่มีคอลัมน์รองรับ — ตัดชั้น Curriculum/PLO ออกเมื่อ 2026-08-04"),
        (None, "", "", "รายวิชาที่ต้องเรียนมาก่อน", "ไม่มี", GAP, "ไม่มีคอลัมน์รองรับ (นอกขอบเขต v1)"),
        (None, "", "", "คำอธิบายรายวิชา", "ศึกษาและปฏิบัติในหัวข้อ หลักการของระบบฐานข้อมูล …",
         GAP, "ไม่มีคอลัมน์รองรับ (นอกขอบเขต v1)"),
    ]),

    ("CourseInstructor", "มอบหมายผู้สอนเข้ารายวิชา — ADMIN เป็นผู้มอบหมาย (FR-22)",
     "1 แถว = 1 คู่ อาจารย์ × รายวิชา", [
        ("id", "VARCHAR(30)", "PK", "รหัสการมอบหมาย", "ci-01", AUTO, None),
        ("courseId", "VARCHAR(30)", "FK · UK", "รายวิชาที่มอบหมาย", COURSE_LINE, IN, "Course.id"),
        ("userId", "VARCHAR(30)", "FK · UK", "อาจารย์ที่ถูกมอบหมาย", "ผศ.วิชัย คงเจริญ", IN, "User.id"),
        ("role", "ENUM(3)", "", "บทบาทในรายวิชา",
         "LEAD (ผู้ประสานงานรายวิชา — มีได้ 1 คนต่อวิชา) · CO · ASSISTANT", IN, None),
        ("assignedAt", "DATETIME(3)", "", "วันที่มอบหมาย", "2025-06-01", AUTO, None),
    ]),

    ("CLO", "ผลลัพธ์การเรียนรู้ระดับรายวิชา — วัดแบบอิงเกณฑ์เสมอ ไม่ขึ้นกับการตัดเกรด",
     "1 แถว = CLO 1 ข้อ", [
        ("id", "VARCHAR(30)", "PK", "รหัส CLO ในระบบ", "clo-01", AUTO, None),
        ("courseId", "VARCHAR(30)", "FK · UK", "รายวิชาที่ CLO สังกัด", COURSE_LINE, IN, "Course.id"),
        ("number", "INT", "UK", "CLO ข้อที่", "1 (ห้ามซ้ำภายในรายวิชาเดียวกัน)", IN, None),
        ("description", "VARCHAR(1000)", "", "ข้อความผลลัพธ์การเรียนรู้",
         "ออกแบบฐานข้อมูลเชิงสัมพันธ์จากโจทย์ที่กำหนดได้", IN, None),
        ("threshold", "DOUBLE", "", "เกณฑ์ผ่านรายคน (%)", "60", IN, None),
        ("bloomLevel", "ENUM(6)", "", "ระดับพฤติกรรมตาม Bloom",
         "สร้างสรรค์ (CREATE) — มาจากคำกริยาที่ขึ้นต้น CLO เอง", IN, None),
        ("classTarget", "DOUBLE", "", "เป้าหมายผู้ผ่านเฉพาะข้อนี้ (%)",
         "60 (เว้นว่าง = ใช้ค่า 70 ของรายวิชา)", IN, None),
    ]),

    ("BehavioralObjective", "จุดประสงค์เชิงพฤติกรรมที่แตกย่อยจาก CLO",
     "1 แถว = 1 จุดประสงค์", [
        ("id", "VARCHAR(30)", "PK", "รหัสจุดประสงค์ในระบบ", "obj-01", AUTO, None),
        ("cloId", "VARCHAR(30)", "FK · UK", "CLO ต้นทาง", "CLO 1 — ออกแบบฐานข้อมูลเชิงสัมพันธ์ฯ",
         IN, "CLO.id"),
        ("number", "INT", "UK", "จุดประสงค์ข้อที่", "1 (ห้ามซ้ำภายใน CLO เดียวกัน)", IN, None),
        ("description", "VARCHAR(1000)", "", "ข้อความจุดประสงค์",
         "เขียนแผนภาพ ER จากความต้องการของผู้ใช้ได้", IN, None),
    ]),

    ("Activity", "กิจกรรมประเมิน — คะแนนถูกบันทึกรายกิจกรรมเท่านั้น",
     "1 แถว = 1 กิจกรรม", [
        ("id", "VARCHAR(30)", "PK", "รหัสกิจกรรมในระบบ", "act-02", AUTO, None),
        ("courseId", "VARCHAR(30)", "FK", "รายวิชาที่กิจกรรมสังกัด", COURSE_LINE, IN, "Course.id"),
        ("name", "VARCHAR(255)", "", "ชื่อกิจกรรม", "สอบกลางภาค", IN, None),
        ("method", "VARCHAR(255)", "", "วิธีการประเมิน", "สอบข้อเขียน", IN, None),
        ("maxScore", "DOUBLE", "", "คะแนนเต็ม", "100 (ต้องมากกว่า 0 — ทุกสูตรหารด้วยค่านี้)", IN, None),
        ("order", "INT", "", "ลำดับการแสดงผล", "2 (ผู้สอนกำหนดเอง จึงต้องเก็บ)", IN, None),
        ("weight", "DOUBLE", "", "สัดส่วนคะแนน (%)", "30 (ทุกกิจกรรมรวมกันได้ 100)", IN, None),
        (None, "", "", "สัปดาห์ที่ประเมิน", "สัปดาห์ที่ 9", GAP,
         "ไม่มีคอลัมน์รองรับ — มคอ.3 หมวดที่ 5 ขอ แต่ v1 ยังไม่เก็บ"),
    ]),

    ("AssessmentCriteria",
     "ผูก Activity กับ CLO พร้อมน้ำหนักต่อคู่ — สิ่งที่ระบบเพิ่มจาก มคอ.3",
     "1 แถว = 1 คู่ กิจกรรม × CLO", [
        ("id", "VARCHAR(30)", "PK", "รหัสการผูกในระบบ", "ac-02", AUTO, None),
        ("activityId", "VARCHAR(30)", "FK · UK", "กิจกรรม", "สอบกลางภาค", IN, "Activity.id"),
        ("cloId", "VARCHAR(30)", "FK · UK", "CLO ที่กิจกรรมนี้วัด",
         "CLO 1 — ออกแบบฐานข้อมูลเชิงสัมพันธ์ฯ", IN, "CLO.id"),
        ("weight", "DOUBLE", "", "น้ำหนักของคู่นี้ (%)",
         "40 (ทุกคู่ของกิจกรรมเดียวกันรวมกันได้ 100)", IN, None),
    ]),

    ("ObjectiveAssessment", "ตามรอยเท่านั้น — ไม่กระทบคะแนน CLO ที่คำนวณได้",
     "1 แถว = 1 คู่ เกณฑ์ × จุดประสงค์", [
        ("id", "VARCHAR(30)", "PK", "รหัสการตามรอยในระบบ", "oa-01", AUTO, None),
        ("criteriaId", "VARCHAR(30)", "FK · UK", "คู่กิจกรรม × CLO ต้นทาง", "สอบกลางภาค × CLO 1",
         IN, "AssessmentCriteria.id"),
        ("objectiveId", "VARCHAR(30)", "FK · UK", "จุดประสงค์เชิงพฤติกรรมที่วัด",
         "เขียนแผนภาพ ER จากความต้องการของผู้ใช้ได้", IN, "BehavioralObjective.id"),
    ]),

    ("Student", "การลงทะเบียน — ไม่ใช่ตารางบุคคล คนเดียวเรียน 5 วิชาได้ 5 แถว",
     "1 แถว = 1 คน 1 วิชา", [
        ("id", "VARCHAR(30)", "PK", "รหัสการลงทะเบียนในระบบ", "std-01", AUTO, None),
        ("studentCode", "VARCHAR(50)", "UK", "รหัสนักศึกษา",
         "65010001 (ห้ามซ้ำภายในรายวิชาเดียวกัน)", IN, None),
        ("name", "VARCHAR(255)", "", "ชื่อ-นามสกุลนักศึกษา", "นายกิตติพงษ์ แสงทอง", IN, None),
        ("courseId", "VARCHAR(30)", "FK · UK", "รายวิชาที่ลงทะเบียน", COURSE_LINE, IN, "Course.id"),
    ]),

    ("Score", "คะแนนดิบรายกิจกรรม — ไม่มีแถว = ยังไม่ประเมิน (ไม่ใช่ 0)",
     "1 แถว = 1 คน 1 กิจกรรม", [
        ("id", "VARCHAR(30)", "PK", "รหัสคะแนนในระบบ", "sc-02", AUTO, None),
        ("studentId", "VARCHAR(30)", "FK · UK", "นักศึกษา", "65010001 นายกิตติพงษ์ แสงทอง",
         IN, "Student.id"),
        ("activityId", "VARCHAR(30)", "FK · UK", "กิจกรรม", "สอบกลางภาค (เต็ม 100)",
         IN, "Activity.id"),
        ("score", "DOUBLE", "", "คะแนนที่ได้", "85", IN, None),
        ("uploadedAt", "DATETIME(3)", "", "วันที่บันทึกคะแนน", "2025-08-01 14:20", AUTO, None),
        (None, "", "", "กรณียังไม่ประเมิน", "เว้นว่างไว้ — ระบบไม่สร้างแถว (ต่างจากการกรอก 0)",
         DERIV, "ความหมายอยู่ที่การไม่มีแถว จึงไม่ต้องมีคอลัมน์"),
    ]),

    ("GradeBand", "ช่วงเกรดของรายวิชา — เรียงจาก minValue และต้องมีช่วงล่างสุดเสมอ",
     "1 แถว = 1 เกรด", [
        ("id", "VARCHAR(30)", "PK", "รหัสช่วงเกรดในระบบ", "gb-01", AUTO, None),
        ("courseId", "VARCHAR(30)", "FK · UK", "รายวิชา", COURSE_LINE, IN, "Course.id"),
        ("grade", "VARCHAR(5)", "UK", "ตัวอักษรเกรด",
         "A (หรือ S / U สำหรับรายวิชาแบบ PASS_FAIL)", IN, None),
        ("minValue", "DOUBLE", "UK", "คะแนนขั้นต่ำของเกรดนี้",
         "80 (อ่านตามวิธีตัดที่รายวิชาเลือกไว้)", IN, None),
        (None, "", "", "ลำดับของเกรด", "A สูงสุด ไล่ลงถึง F", DERIV,
         "เรียงจาก minValue มากไปน้อย จึงไม่ต้องเก็บคอลัมน์ order"),
    ]),

    ("StudentGrade", "เกรดสุดท้ายรายคน — ปรับมือได้แต่ต้องระบุเหตุผล",
     "1 แถว = 1 นักศึกษา 1 รายวิชา", [
        ("id", "VARCHAR(30)", "PK", "รหัสเกรดในระบบ", "sg-01", AUTO, None),
        ("studentId", "VARCHAR(30)", "FK · UK", "นักศึกษา",
         "65010001 นายกิตติพงษ์ แสงทอง (รายวิชาติดมากับนักศึกษาอยู่แล้ว)", AUTO, "Student.id"),
        ("totalPercent", "DOUBLE", "", "คะแนนรวมถ่วงน้ำหนัก (0-100)",
         "83.60 (บันทึกค่าจริงไว้ตอนตัดเกรด)", AUTO, None),
        ("grade", "VARCHAR(5)", "", "เกรดที่ได้", "A", AUTO, None),
        ("overrideReason", "VARCHAR(1000)", "", "เหตุผลการปรับเกรด",
         "เว้นว่าง = ใช้เกรดที่ระบบคำนวณ · กรอกแล้ว = ปรับมือ (ต้องมีเหตุผลเสมอ)", IN, None),
        (None, "", "", "คะแนนมาตรฐาน T / Z และค่าสถิติของห้อง",
         "อิงกลุ่มเท่านั้น — คำนวณตอนกดตัดเกรด แล้วไม่เก็บ", DERIV,
         "ผลที่เก็บคือตัวเกรด ไม่ใช่ที่มาของเกรด"),
    ]),

    ("ScoreUploadLog", "ตามรอยการนำเข้าคะแนน — อยู่นอกเส้นทางการคำนวณ",
     "1 แถว = 1 ครั้งที่อัปโหลด", [
        ("id", "VARCHAR(30)", "PK", "รหัสบันทึกในระบบ", "log-01", AUTO, None),
        ("courseId", "VARCHAR(30)", "FK", "รายวิชา", COURSE_LINE, AUTO, "Course.id"),
        ("uploadedBy", "VARCHAR(30)", "FK", "ผู้อัปโหลด", "ผศ.วิชัย คงเจริญ", AUTO, "User.id"),
        ("fileName", "VARCHAR(255)", "", "ชื่อไฟล์ที่อัปโหลด", "scores_midterm_03376120.xlsx", IN, None),
        ("recordsOk", "INT", "", "จำนวนแถวที่นำเข้าสำเร็จ", "7", AUTO, None),
        ("recordsFail", "INT", "", "จำนวนแถวที่ผิดพลาด", "1", AUTO, None),
        ("createdAt", "DATETIME(3)", "", "เวลาที่อัปโหลด", "2025-08-01 14:20", AUTO, None),
    ]),
]


# ===========================================================================
# SELF-CHECK against the file MySQL Workbench actually reads
# ===========================================================================
def parse_sql():
    text = SQL.read_text(encoding="utf-8").split("SET FOREIGN_KEY_CHECKS = 0;")[1]
    out = {}
    for m in re.finditer(r"CREATE TABLE `?(\w+)`?\s*\((.*?)\n\) ENGINE", text, re.S):
        name, inner = m.group(1), m.group(2)
        cols = re.findall(
            r"^\s*`?(\w+)`?\s+(VARCHAR\(\d+\)|INT|DOUBLE|DECIMAL\(\d+,\d+\)|"
            r"DATETIME\(\d\)|TINYINT\(\d\)|ENUM\([^)]*\))",
            inner, re.M)
        uq = set()
        for group in re.findall(r"UNIQUE KEY \w+ \(([^)]+)\)", inner):
            uq |= {c.strip().strip("`") for c in group.split(",")}
        pk = {c.strip().strip("`")
              for c in re.search(r"PRIMARY KEY \(([^)]+)\)", inner).group(1).split(",")}
        fk = {m2.group(1): f"{m2.group(2)}.{m2.group(3)}" for m2 in re.finditer(
            r"FOREIGN KEY \((\w+)\) REFERENCES `?(\w+)`? \((\w+)\)", inner)}
        out[name] = {"cols": [c[0] for c in cols],
                     "types": {c[0]: c[1] for c in cols},
                     "pk": pk, "uq": uq, "fk": fk}
    return out


def norm_type(t):
    """ENUM(...) is written as ENUM(n) in the workbook — compare the count.

    The SQL wraps long ENUM lists over several lines, so collapse whitespace
    before matching; `.` does not cross a newline.
    """
    t = " ".join(t.split())
    m = re.match(r"ENUM\((.*)\)$", t)
    if m:
        return f"ENUM({len(m.group(1).split(','))})"
    return t


def self_check(sql):
    errs = []
    if len(sql) != len(TABLES):
        errs.append(f"SQL has {len(sql)} tables, workbook defines {len(TABLES)}")
    for name, _, _, rows in TABLES:
        if name not in sql:
            errs.append(f"{name}: not in the SQL at all")
            continue
        t = sql[name]
        mine = [r for r in rows if r[0] is not None]
        if [r[0] for r in mine] != t["cols"]:
            errs.append(f"{name}: columns/order differ\n"
                        f"    workbook {[r[0] for r in mine]}\n"
                        f"    sql      {t['cols']}")
            continue
        for col, typ, key, *_ in mine:
            if norm_type(t["types"][col]) != typ:
                errs.append(f"{name}.{col}: type {typ} != {t['types'][col]}")
            want = set()
            if col in t["pk"]:
                want.add("PK")
            if col in t["fk"]:
                want.add("FK")
            if col in t["uq"]:
                want.add("UK")
            got = {k.strip() for k in key.split("·") if k.strip()}
            if want != got:
                errs.append(f"{name}.{col}: key '{key or '—'}' != {sorted(want) or '—'}")
    for name in sql:
        if name not in [t[0] for t in TABLES]:
            errs.append(f"{name}: in the SQL but has no sheet")
    return errs


SQLDEF = parse_sql()
errors = self_check(SQLDEF)
if errors:
    print("BUILD STOPPED — workbook does not match cmas_app_mysql_v4.sql:\n")
    print("\n".join("  - " + e for e in errors))
    sys.exit(1)

FK_MAP = [(name, col, ref)
          for name, _, _, rows in TABLES
          for col, _, _, _, _, _, ref in rows
          if ref and "." in str(ref) and col]

# ===========================================================================
# RENDER
# ===========================================================================
wb = Workbook()


def set_widths(ws, widths):
    for c, w in widths.items():
        ws.column_dimensions[c].width = w


def header_row(ws, row, headers, start_col=2, fill=FILL_UI):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=start_col + i, value=h)
        c.font = F_HEAD
        c.fill = fill
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 24


def cell(ws, row, col, value, *, zebra=False, font=None, fill=None, align=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = font or F_BODY
    c.border = BORDER
    c.fill = fill or (FILL_ROW if zebra else FILL_WHITE)
    c.alignment = align or Alignment(vertical="center", wrap_text=True)
    return c


def key_cell(ws, row, col, key):
    fg, bg = KEY_STYLE.get(key, KEY_STYLE[""])
    c = ws.cell(row=row, column=col, value=key or "—")
    c.font = Font(name="Calibri", size=9, bold=bool(key), color=fg)
    c.fill = PatternFill("solid", fgColor=bg)
    c.border = BORDER
    c.alignment = Alignment(horizontal="center", vertical="center")
    return c


# --------------------------------------------------------------------- index
ws = wb.active
ws.title = "สารบัญ"
ws.sheet_view.showGridLines = False
set_widths(ws, {"A": 3, "B": 24, "C": 46, "D": 11, "E": 8, "F": 8, "G": 10})

ws["A1"] = "ตารางทั้งหมดในแผนภาพ ER"
ws["A1"].font = F_TITLE
ws["A2"] = (f"ตรงกับ docs/reference/db/mysql/cmas_app_mysql_v4.sql ซึ่งเป็นไฟล์ที่ MySQL Workbench "
            f"ใช้สร้างแผนภาพ · สคริปต์จะไม่ยอมสร้างไฟล์นี้ถ้าสองฝั่งไม่ตรงกัน")
ws["A2"].font = F_SUB

r = 4
header_row(ws, r, ["ตาราง", "เก็บอะไร", "คอลัมน์", "PK", "FK", "UNIQUE"])
r += 1
for i, (name, caption, grain, rows) in enumerate(TABLES):
    z = i % 2 == 1
    t = SQLDEF[name]
    cell(ws, r, 2, name, zebra=z, font=F_LABEL)
    cell(ws, r, 3, caption, zebra=z)
    cell(ws, r, 4, len(t["cols"]), zebra=z, align=Alignment(horizontal="center", vertical="center"))
    cell(ws, r, 5, len(t["pk"]), zebra=z, align=Alignment(horizontal="center", vertical="center"),
         font=Font(name="Calibri", size=10, bold=True, color=C_CALC))
    cell(ws, r, 6, len(t["fk"]) or "—", zebra=z,
         align=Alignment(horizontal="center", vertical="center"),
         font=Font(name="Calibri", size=10, bold=True, color=C_UI))
    cell(ws, r, 7, len(t["uq"]) or "—", zebra=z,
         align=Alignment(horizontal="center", vertical="center"),
         font=Font(name="Calibri", size=10, bold=True, color=C_TQF))
    r += 1

cell(ws, r, 2, "รวม", font=F_LABEL, fill=FILL_ROW)
cell(ws, r, 3, f"{len(TABLES)} ตาราง", fill=FILL_ROW, font=F_LABEL)
cell(ws, r, 4, sum(len(t["cols"]) for t in SQLDEF.values()), fill=FILL_ROW, font=F_LABEL,
     align=Alignment(horizontal="center", vertical="center"))
cell(ws, r, 5, sum(len(t["pk"]) for t in SQLDEF.values()), fill=FILL_ROW, font=F_LABEL,
     align=Alignment(horizontal="center", vertical="center"))
cell(ws, r, 6, sum(len(t["fk"]) for t in SQLDEF.values()), fill=FILL_ROW, font=F_LABEL,
     align=Alignment(horizontal="center", vertical="center"))
cell(ws, r, 7, "", fill=FILL_ROW)
r += 2

ws.cell(row=r, column=2, value="ความสัมพันธ์ทั้งหมด (Foreign Key)").font = F_H2
r += 1
header_row(ws, r, ["จากตาราง", "คอลัมน์", "ชี้ไปที่"])
r += 1
for i, (tbl, col, ref) in enumerate(FK_MAP):
    z = i % 2 == 1
    cell(ws, r, 2, tbl, zebra=z, font=F_LABEL)
    cell(ws, r, 3, col, zebra=z, font=F_MONO)
    cell(ws, r, 4, "→  " + ref, zebra=z, font=F_MONO)
    r += 1

r += 1
ws.cell(row=r, column=2, value="ความหมายของสี").font = F_H2
r += 1
header_row(ws, r, ["สัญลักษณ์", "หมายความว่า"])
r += 1
legend = [
    ("PK", "Primary Key — คีย์หลักของตาราง ทุกตารางใช้ id ที่ระบบสร้าง (cuid)", KEY_STYLE["PK"]),
    ("FK", "Foreign Key — ชี้ไปยังแถวในอีกตารางหนึ่ง", KEY_STYLE["FK"]),
    ("UK", "Unique Key — คีย์ธรรมชาติที่กันข้อมูลซ้ำ (เช่น รหัสวิชา + ภาค + ปี + หมู่)",
     KEY_STYLE["UK"]),
    ("FK · UK", "เป็นทั้งสองอย่าง — ชี้ไปตารางอื่น และเป็นส่วนหนึ่งของคีย์กันซ้ำ", KEY_STYLE["FK"]),
]
for k, meaning, (fg, bg) in legend:
    c = ws.cell(row=r, column=2, value=k)
    c.font = Font(name="Calibri", size=9, bold=True, color=fg)
    c.fill = PatternFill("solid", fgColor=bg)
    c.border = BORDER
    c.alignment = Alignment(horizontal="center", vertical="center")
    cell(ws, r, 3, meaning)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
    r += 1

rows_legend = [
    ("ฟ้าอ่อน", "ผู้ใช้กรอกค่านี้เอง", FILL_INPUT, F_BODY),
    ("เขียวอ่อน", "ระบบสร้างหรือคำนวณให้อัตโนมัติ", FILL_AUTO, F_BODY),
    ("เทา", "หน้าจอแสดงค่านี้ได้ แต่ไม่มีคอลัมน์เก็บ — คำนวณจากคอลัมน์ที่มีอยู่แล้ว",
     FILL_DERIV, F_BODY),
    ("ส้ม", "มคอ.3 ขอข้อมูลนี้ แต่ v1 ไม่มีคอลัมน์รองรับ — ช่องว่างของขอบเขต", FILL_GAP, F_GAP),
]
for name, meaning, fill, font in rows_legend:
    c = cell(ws, r, 2, name, fill=fill, font=font,
             align=Alignment(horizontal="center", vertical="center"))
    cell(ws, r, 3, meaning)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
    r += 1

ws.freeze_panes = "B5"

# ---------------------------------------------------------------- per table
for name, caption, grain, rows in TABLES:
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    set_widths(ws, {"A": 3, "B": 10, "C": 34, "D": 56, "E": 44, "F": 16})

    ws["A1"] = f"ตาราง {name}"
    ws["A1"].font = F_TITLE
    ws["A2"] = f"{caption}  ·  {grain}"
    ws["A2"].font = F_SUB

    t = SQLDEF[name]
    ins = [f"{a} → {b}" for a, b in
           [(f"{tb}.{c}", name) for tb, c, ref in FK_MAP if ref.split(".")[0] == name]]
    outs = [f"{c} → {ref}" for tb, c, ref in FK_MAP if tb == name]
    rel = []
    if outs:
        rel.append("ชี้ออกไป: " + " · ".join(outs))
    if ins:
        rel.append("ถูกชี้เข้ามา: " + " · ".join(x.split(" → ")[0] for x in ins))
    c = ws.cell(row=3, column=2, value=("  |  ".join(rel)) if rel else
                "ไม่มีความสัมพันธ์ออกหรือเข้า")
    c.font = F_MONO
    ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=6)

    header_row(ws, 5, ["คีย์", "หัวข้อ", "ค่าที่กรอกในระบบ", "เก็บที่คอลัมน์", "ชนิดข้อมูล"])
    r = 6
    for i, (col, typ, key, label, example, kind, ref) in enumerate(rows):
        z = i % 2 == 1
        if col is None:
            key_cell(ws, r, 2, "")
            cell(ws, r, 3, label, zebra=z, font=F_LABEL)
            cell(ws, r, 4, example, zebra=z,
                 fill=FILL_GAP if kind == GAP else FILL_DERIV)
            cell(ws, r, 5, ("— ไม่มีคอลัมน์ · " if kind == GAP else "— ไม่เก็บ · ") + str(ref),
                 zebra=z, font=F_GAP if kind == GAP else F_DERIVF,
                 fill=FILL_GAP if kind == GAP else FILL_DERIV)
            cell(ws, r, 6, "—", zebra=z, font=F_TYPE,
                 align=Alignment(horizontal="center", vertical="center"))
        else:
            key_cell(ws, r, 2, key)
            cell(ws, r, 3, label, zebra=z, font=F_LABEL)
            cell(ws, r, 4, example, zebra=z,
                 fill=FILL_AUTO if kind == AUTO else FILL_INPUT)
            target = f"{name}.{col}" + (f"  →  {ref}" if ref else "")
            cell(ws, r, 5, target, zebra=z, font=F_MONO)
            cell(ws, r, 6, typ, zebra=z, font=F_TYPE,
                 align=Alignment(horizontal="center", vertical="center"))
        r += 1

    ws.freeze_panes = "C6"

wb.save(OUT)
print(f"wrote {OUT}")
print(f"{len(TABLES)} sheets · {sum(len(t['cols']) for t in SQLDEF.values())} columns · "
      f"{sum(len(t['pk']) for t in SQLDEF.values())} PK · {len(FK_MAP)} FK · "
      f"verified against {SQL.name}")
