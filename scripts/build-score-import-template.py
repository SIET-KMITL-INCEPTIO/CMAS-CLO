"""
Build docs/excel/CMAS-Score-Import-Template.xlsx — the REAL, blank, protected
Excel workbook an instructor fills in by hand: student roster (FR-50/51/52)
and score entry (FR-60/61/62/66). Unlike the ER-Diagram workbook (which is
documentation with fictional example data), every input cell here is genuinely
empty and ready for real numbers — headers and formulas are locked so the
sheet's structure survives, only the actual data cells are editable.

Run: python scripts/build-score-import-template.py
Output: docs/excel/CMAS-Score-Import-Template.xlsx
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

# ---------------------------------------------------------------------------
# Palette (same family as the other CMAS docs) + form-specific accents
# ---------------------------------------------------------------------------
C_PRIMARY = "19151B"
C_SECONDARY = "5C5C78"
C_SECONDARY_CONTAINER = "DEDDFE"
C_ON_SURFACE = "1A1C1D"
C_ON_SURFACE_VARIANT = "4A454A"
C_OUTLINE = "CCC4CA"

C_LOCKED_BG = "EEEEF0"      # structure the instructor must not edit
C_EDIT_BG = "FFFFFF"        # real input cell
C_EDIT_BORDER = "1A56C4"    # blue border marks "type here"
C_WARN_BG = "FEF7E0"
C_WARN_TEXT = "B06000"
C_DANGER_BG = "FCE8E6"
C_DANGER_TEXT = "C5221F"
C_OK_BG = "E6F4EA"
C_OK_TEXT = "137333"

thin = Side(style="thin", color=C_OUTLINE)
BORDER_ALL = Border(left=thin, right=thin, top=thin, bottom=thin)
edit_side = Side(style="thin", color=C_EDIT_BORDER)
BORDER_EDIT = Border(left=edit_side, right=edit_side, top=edit_side, bottom=edit_side)

FONT_TITLE = Font(name="Calibri", size=16, bold=True, color=C_PRIMARY)
FONT_SUBTITLE = Font(name="Calibri", size=10.5, italic=True, color=C_ON_SURFACE_VARIANT)
FONT_H2 = Font(name="Calibri", size=12, bold=True, color=C_PRIMARY)
FONT_LABEL = Font(name="Calibri", size=10, bold=True, color=C_ON_SURFACE)
FONT_BODY = Font(name="Calibri", size=10, color=C_ON_SURFACE)
FONT_MUTED = Font(name="Calibri", size=9.5, color=C_ON_SURFACE_VARIANT)
FONT_HEADER = Font(name="Calibri", size=10, bold=True, color="FFFFFF")

FILL_HEADER = PatternFill("solid", fgColor=C_PRIMARY)
FILL_SECONDARY = PatternFill("solid", fgColor=C_SECONDARY_CONTAINER)
FILL_LOCKED = PatternFill("solid", fgColor=C_LOCKED_BG)
FILL_EDIT = PatternFill("solid", fgColor=C_EDIT_BG)
FILL_WARN = PatternFill("solid", fgColor=C_WARN_BG)
FILL_DANGER = PatternFill("solid", fgColor=C_DANGER_BG)
FILL_OK = PatternFill("solid", fgColor=C_OK_BG)

LOCK = Protection(locked=True)
UNLOCK = Protection(locked=False)

N_STUDENT_ROWS = 40
N_ACTIVITY_COLS = 8  # D..K

wb = Workbook()

# ===========================================================================
# 0. คำแนะนำ (Instructions)
# ===========================================================================
ws = wb.active
ws.title = "คำแนะนำ"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 100

ws["B2"] = "CMAS — แบบฟอร์มนำเข้ารายชื่อนักศึกษาและคะแนน (ฉบับกรอกจริง)"
ws["B2"].font = FONT_TITLE
ws["B3"] = "ไฟล์นี้ไม่มีข้อมูลจำลอง — ทุกช่องพร้อมให้อาจารย์กรอกข้อมูลจริงได้ทันที"
ws["B3"].font = FONT_SUBTITLE

r = 5
def h2(text):
    global r
    c = ws.cell(row=r, column=2, value=text)
    c.font = FONT_H2
    r += 1

def body(text, gap=1):
    global r
    c = ws.cell(row=r, column=2, value=text)
    c.font = FONT_BODY
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 30
    r += gap

h2("ลำดับขั้นตอนการใช้งาน")
steps = [
    "1. เปิด sheet \"01 รายชื่อนักศึกษา\" กรอกข้อมูลรายวิชาที่ด้านบน แล้วกรอกรหัสนักศึกษาและชื่อ-นามสกุลทีละแถว (ไม่ต้องเรียงลำดับ ไม่ต้องกรอกครบทุกแถวถ้านักศึกษาน้อยกว่า 40 คน)",
    "2. เปิด sheet \"02 กรอกคะแนน\" — รหัสนักศึกษาและชื่อจะดึงมาจาก sheet รายชื่อนักศึกษาให้อัตโนมัติ ไม่ต้องพิมพ์ซ้ำ",
    "3. กรอกชื่อกิจกรรมและคะแนนเต็มในแถว \"ตั้งค่ากิจกรรม\" ก่อน (สูงสุด 8 กิจกรรม) — หัวคอลัมน์คะแนนจะอัปเดตชื่อและคะแนนเต็มให้อัตโนมัติ",
    "4. กรอกคะแนนแต่ละคนในแต่ละกิจกรรม — ถ้ายังไม่มีคะแนน ให้ \"เว้นว่างไว้\" ห้ามใส่ 0 เพราะ 0 หมายถึง \"สอบแล้วได้ 0 คะแนนจริง\" ส่วนช่องว่างหมายถึง \"ยังไม่ได้ประเมิน\" ทั้งสองอย่างมีผลต่อการคำนวณต่างกัน",
    "5. บันทึกไฟล์ (Save) แล้วอัปโหลดกลับเข้าระบบที่หน้า \"นำเข้าคะแนน\" — ระบบจะแสดงตัวอย่าง (preview) ให้ตรวจสอบก่อนบันทึกจริงเสมอ ไม่มีการบันทึกทันทีโดยไม่ให้ตรวจ",
]
for s in steps:
    body(s)

r += 1
h2("กฎที่ระบบจะตรวจสอบอัตโนมัติตอนอัปโหลด")
rules = [
    "รหัสนักศึกษาที่กรอกต้องมีอยู่ในรายวิชานี้แล้ว (เพิ่มที่ sheet รายชื่อนักศึกษาก่อน หรือในระบบก่อนอัปโหลด)",
    "คะแนนต้องเป็นตัวเลขเท่านั้น (ไม่ใช่ตัวอักษรหรือเครื่องหมาย เช่น \"-\", \"ขาดสอบ\")",
    "คะแนนต้องไม่เกินคะแนนเต็มของกิจกรรมนั้น และต้องไม่ติดลบ",
    "รหัสนักศึกษาต้องไม่ซ้ำกันภายในไฟล์เดียวกัน",
    "ชื่อคอลัมน์กิจกรรมในไฟล์ต้องตรงกับกิจกรรมที่ตั้งไว้ในระบบ",
    "แถวที่ผิดพลาดจะไม่ทำให้ทั้งไฟล์ล้มเหลว — ระบบจะบันทึกเฉพาะแถวที่ถูกต้อง และให้ดาวน์โหลดรายการแถวที่ผิดพลาดพร้อมเหตุผลได้",
]
for rule in rules:
    body("•  " + rule)

r += 1
h2("สีของช่องกรอก (Legend)")
legend = [
    (FILL_LOCKED, "สีเทา — โครงสร้างของฟอร์ม ห้ามแก้ไข/ลบ (ชื่อคอลัมน์, ลำดับ)"),
    (FILL_EDIT, "สีขาวขอบน้ำเงิน — ช่องสำหรับกรอกข้อมูลจริง"),
    (FILL_WARN, "สีเหลืองอ่อน — ข้อควรระวัง"),
]
for fill, text in legend:
    c1 = ws.cell(row=r, column=2)
    c1.fill = fill
    c1.border = BORDER_ALL
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=2)
    ws.row_dimensions[r].height = 18
    c2 = ws.cell(row=r, column=3, value=text)
    c2.font = FONT_BODY
    r += 1

r += 2
h2("ตัวอย่างประกอบ (ภาพประกอบเท่านั้น — ไม่ใช่ข้อมูลจริง ห้ามคัดลอกไปใช้)")
r += 1
ex_headers = ["รหัสนักศึกษา", "ชื่อ-นามสกุล", "Quiz 1 (เต็ม 10)", "ความหมาย"]
for i, htext in enumerate(ex_headers, start=2):
    c = ws.cell(row=r, column=i, value=htext)
    c.font = FONT_HEADER
    c.fill = FILL_HEADER
    c.border = BORDER_ALL
    c.alignment = Alignment(horizontal="center", wrap_text=True)
ws.column_dimensions["C"].width = 22
ws.column_dimensions["D"].width = 16
ws.column_dimensions["E"].width = 34
r += 1
example_rows = [
    ("65010099", "นายตัวอย่าง หนึ่ง", 8, "สอบแล้ว ได้ 8 คะแนนจาก 10 — กรอกตัวเลขตามจริง"),
    ("65010098", "นายตัวอย่าง สอง", 0, "สอบแล้ว ได้ 0 คะแนนจริง — พิมพ์ 0 ได้ตามจริง"),
    ("65010097", "นายตัวอย่าง สาม", "(เว้นว่าง)", "ยังไม่ได้สอบ/ยังไม่ให้คะแนน — ห้ามใส่ 0 ให้เว้นว่างไว้เท่านั้น"),
]
for code, name, score, meaning in example_rows:
    vals = [code, name, score, meaning]
    for i, v in enumerate(vals, start=2):
        c = ws.cell(row=r, column=i, value=v)
        c.font = FONT_BODY
        c.border = BORDER_ALL
        c.alignment = Alignment(vertical="center", wrap_text=(i == 5))
    r += 1

for row in ws.iter_rows(min_row=1, max_row=r, min_col=1, max_col=6):
    for cell in row:
        cell.protection = LOCK
ws.protection.sheet = True

# ===========================================================================
# Shared: course-info block builder
# ===========================================================================

def course_info_block(ws: Worksheet, start_row: int, sheet_label: str):
    ws.cell(row=start_row, column=1, value=sheet_label).font = FONT_H2
    fields = [
        ("รหัสวิชา", "B", "C"), ("ชื่อวิชา", "D", "F"),
        ("ภาคการศึกษา", "B", "C"), ("ปีการศึกษา (พ.ศ.)", "D", "F"),
        ("กลุ่มเรียน (Section)", "B", "C"), ("ผู้สอน", "D", "F"),
    ]
    r0 = start_row + 1
    coords = {}
    for i in range(3):
        row = r0 + i
        label1, c1, c1e = fields[i * 2]
        label2, c2, c2e = fields[i * 2 + 1]
        l1 = ws.cell(row=row, column=1, value=label1 + ":")
        l1.font = FONT_LABEL
        v1 = ws[f"{c1}{row}"]
        ws.merge_cells(f"{c1}{row}:{c1e}{row}")
        v1.fill = FILL_EDIT
        v1.border = BORDER_EDIT
        v1.protection = UNLOCK
        coords[label1] = f"{c1}{row}"

        l2 = ws.cell(row=row, column=7, value=label2 + ":")
        l2.font = FONT_LABEL
        v2 = ws[f"{c2 if False else 'H'}{row}"]
        ws.merge_cells(f"H{row}:J{row}")
        v2.fill = FILL_EDIT
        v2.border = BORDER_EDIT
        v2.protection = UNLOCK
        coords[label2] = f"H{row}"
        ws.row_dimensions[row].height = 18
    return r0 + 3  # next free row after block


# ===========================================================================
# 1. 01 รายชื่อนักศึกษา  (Student Roster — FR-50/51/52)
# ===========================================================================
ws = wb.create_sheet("01 รายชื่อนักศึกษา")
ws.sheet_view.showGridLines = False
widths = {"A": 6, "B": 18, "C": 30, "D": 18, "E": 6, "F": 12, "G": 20, "H": 18, "I": 6, "J": 18}
for col, w in widths.items():
    ws.column_dimensions[col].width = w

ws["A1"] = "รายชื่อนักศึกษาในรายวิชา"
ws["A1"].font = FONT_TITLE
ws.merge_cells("A1:F1")

next_row = course_info_block(ws, 3, "ข้อมูลรายวิชา")

header_row = next_row + 1
headers = ["ลำดับ", "รหัสนักศึกษา", "ชื่อ-นามสกุล"]
for i, htext in enumerate(headers, start=1):
    c = ws.cell(row=header_row, column=i, value=htext)
    c.font = FONT_HEADER
    c.fill = FILL_HEADER
    c.border = BORDER_ALL
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.protection = LOCK
ws.row_dimensions[header_row].height = 20

first_data_row = header_row + 1
for i in range(N_STUDENT_ROWS):
    row = first_data_row + i
    c_no = ws.cell(row=row, column=1, value=i + 1)
    c_no.font = FONT_MUTED
    c_no.fill = FILL_LOCKED
    c_no.border = BORDER_ALL
    c_no.alignment = Alignment(horizontal="center")
    c_no.protection = LOCK

    c_code = ws.cell(row=row, column=2)
    c_code.fill = FILL_EDIT
    c_code.border = BORDER_EDIT
    c_code.protection = UNLOCK
    c_code.number_format = "@"  # text, so leading zeros in studentCode survive

    c_name = ws.cell(row=row, column=3)
    c_name.fill = FILL_EDIT
    c_name.border = BORDER_EDIT
    c_name.protection = UNLOCK

last_data_row = first_data_row + N_STUDENT_ROWS - 1
ws.freeze_panes = f"B{first_data_row}"

note_row = last_data_row + 2
note = ws.cell(row=note_row, column=1, value="รหัสต้องไม่ซ้ำกันภายในวิชานี้ (FR-51) — เว้นแถวว่างได้ถ้านักศึกษาน้อยกว่า 40 คน ไม่ต้องลบแถวที่เหลือ")
note.font = FONT_MUTED
ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=6)

for row in ws.iter_rows(min_row=1, max_row=note_row + 1, min_col=1, max_col=10):
    for cell in row:
        if cell.protection is None or cell.protection.locked is None:
            cell.protection = LOCK
ws.protection.sheet = True

ROSTER_SHEET_NAME = "01 รายชื่อนักศึกษา"
ROSTER_FIRST_DATA_ROW = first_data_row

# ===========================================================================
# 2. 02 กรอกคะแนน  (Score Entry — FR-60/61/62/66)
# ===========================================================================
ws = wb.create_sheet("02 กรอกคะแนน")
ws.sheet_view.showGridLines = False
widths = {"A": 6, "B": 18, "C": 26}
for col, w in widths.items():
    ws.column_dimensions[col].width = w
for i in range(N_ACTIVITY_COLS):
    ws.column_dimensions[get_column_letter(4 + i)].width = 16

ws["A1"] = "กรอกคะแนนรายกิจกรรม"
ws["A1"].font = FONT_TITLE
ws.merge_cells("A1:F1")

next_row = course_info_block(ws, 3, "ข้อมูลรายวิชา")

# --- Activity setup block ---
setup_label_row = next_row + 1
c = ws.cell(row=setup_label_row, column=1, value="ตั้งค่ากิจกรรมการประเมิน (Activity) — กรอกชื่อกิจกรรมและคะแนนเต็มที่นี่ก่อน สูงสุด 8 กิจกรรม")
c.font = FONT_LABEL
ws.merge_cells(start_row=setup_label_row, start_column=1, end_row=setup_label_row, end_column=3 + N_ACTIVITY_COLS)

name_row = setup_label_row + 1
max_row_ = setup_label_row + 2
ws.cell(row=name_row, column=3, value="ชื่อกิจกรรม →").font = FONT_MUTED
ws.cell(row=max_row_, column=3, value="คะแนนเต็ม →").font = FONT_MUTED
for c_ in (ws.cell(row=name_row, column=3), ws.cell(row=max_row_, column=3)):
    c_.protection = LOCK
    c_.alignment = Alignment(horizontal="right")

activity_name_cells = []
activity_max_cells = []
for i in range(N_ACTIVITY_COLS):
    col = 4 + i
    letter = get_column_letter(col)
    nc = ws.cell(row=name_row, column=col)
    nc.fill = FILL_EDIT
    nc.border = BORDER_EDIT
    nc.protection = UNLOCK
    nc.alignment = Alignment(horizontal="center")
    activity_name_cells.append(f"{letter}{name_row}")

    mc = ws.cell(row=max_row_, column=col)
    mc.fill = FILL_EDIT
    mc.border = BORDER_EDIT
    mc.protection = UNLOCK
    mc.alignment = Alignment(horizontal="center")
    mc.number_format = "0.##"
    activity_max_cells.append(f"{letter}{max_row_}")
ws.row_dimensions[name_row].height = 18
ws.row_dimensions[max_row_].height = 18

warn_row = max_row_ + 1
wcell = ws.cell(row=warn_row, column=1, value="ชื่อกิจกรรมต้องตรงกับที่ตั้งไว้ในระบบทุกตัวอักษร มิฉะนั้นระบบจะปฏิเสธคอลัมน์นั้นทั้งคอลัมน์ตอนอัปโหลด (FR-69)")
wcell.font = Font(size=9.5, italic=True, color=C_WARN_TEXT)
wcell.fill = FILL_WARN
ws.merge_cells(start_row=warn_row, start_column=1, end_row=warn_row, end_column=3 + N_ACTIVITY_COLS)
wcell.protection = LOCK

# --- score table header (dynamic — reflects activity setup above) ---
header_row = warn_row + 2
h1 = ws.cell(row=header_row, column=1, value="ลำดับ")
h2c = ws.cell(row=header_row, column=2, value="รหัสนักศึกษา")
h3 = ws.cell(row=header_row, column=3, value="ชื่อ-นามสกุล")
for hc in (h1, h2c, h3):
    hc.font = FONT_HEADER
    hc.fill = FILL_HEADER
    hc.border = BORDER_ALL
    hc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    hc.protection = LOCK

for i in range(N_ACTIVITY_COLS):
    col = 4 + i
    letter = get_column_letter(col)
    name_ref = activity_name_cells[i]
    max_ref = activity_max_cells[i]
    formula = f'=IF({name_ref}="","กิจกรรมที่ {i + 1}",{name_ref}&" (เต็ม "&IF({max_ref}="","?",{max_ref})&")")'
    hc = ws.cell(row=header_row, column=col, value=formula)
    hc.font = FONT_HEADER
    hc.fill = FILL_HEADER
    hc.border = BORDER_ALL
    hc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    hc.protection = LOCK
ws.row_dimensions[header_row].height = 30

first_data_row = header_row + 1
roster_offset = ROSTER_FIRST_DATA_ROW - first_data_row  # to map score-row -> matching roster-row

dvs = []
for i in range(N_ACTIVITY_COLS):
    col = 4 + i
    letter = get_column_letter(col)
    max_ref = activity_max_cells[i]
    dv = DataValidation(
        type="decimal",
        operator="between",
        formula1="0",
        formula2=f'=IF({max_ref}="",100,{max_ref})',
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="คะแนนไม่ถูกต้อง",
        error="คะแนนต้องเป็นตัวเลข ตั้งแต่ 0 ถึงคะแนนเต็มของกิจกรรมนี้ — เว้นว่างไว้ถ้ายังไม่ได้ประเมิน",
        promptTitle="กรอกคะแนน",
        prompt="ใส่ตัวเลข 0 ถึงคะแนนเต็ม หรือเว้นว่างไว้ถ้ายังไม่ประเมิน (ห้ามใส่ 0 แทนการยังไม่ประเมิน)",
        showInputMessage=True,
    )
    ws.add_data_validation(dv)
    dvs.append((dv, letter))

for i in range(N_STUDENT_ROWS):
    row = first_data_row + i
    roster_row = row + roster_offset

    c_no = ws.cell(row=row, column=1, value=i + 1)
    c_no.font = FONT_MUTED
    c_no.fill = FILL_LOCKED
    c_no.border = BORDER_ALL
    c_no.alignment = Alignment(horizontal="center")
    c_no.protection = LOCK

    code_formula = f"=IF('{ROSTER_SHEET_NAME}'!B{roster_row}=\"\",\"\",'{ROSTER_SHEET_NAME}'!B{roster_row})"
    c_code = ws.cell(row=row, column=2, value=code_formula)
    c_code.fill = FILL_LOCKED
    c_code.border = BORDER_ALL
    c_code.protection = LOCK
    c_code.number_format = "@"

    name_formula = f"=IF('{ROSTER_SHEET_NAME}'!C{roster_row}=\"\",\"\",'{ROSTER_SHEET_NAME}'!C{roster_row})"
    c_name = ws.cell(row=row, column=3, value=name_formula)
    c_name.fill = FILL_LOCKED
    c_name.border = BORDER_ALL
    c_name.protection = LOCK

    for j in range(N_ACTIVITY_COLS):
        col = 4 + j
        sc = ws.cell(row=row, column=col)
        sc.fill = FILL_EDIT
        sc.border = BORDER_EDIT
        sc.protection = UNLOCK
        dv, letter = dvs[j]
        dv.add(sc)

last_data_row = first_data_row + N_STUDENT_ROWS - 1
ws.freeze_panes = f"D{first_data_row}"

note_row = last_data_row + 2
note = ws.cell(
    row=note_row, column=1,
    value="รหัสนักศึกษา/ชื่อ ดึงมาจาก sheet \"01 รายชื่อนักศึกษา\" อัตโนมัติ — ถ้าเพิ่ม/แก้ชื่อ ให้แก้ที่ sheet นั้น ไม่ต้องแก้ที่นี่",
)
note.font = FONT_MUTED
ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=3 + N_ACTIVITY_COLS)
note.protection = LOCK

for row in ws.iter_rows(min_row=1, max_row=note_row + 1, min_col=1, max_col=3 + N_ACTIVITY_COLS):
    for cell in row:
        if cell.protection is None or cell.protection.locked is None:
            cell.protection = LOCK
ws.protection.sheet = True

wb.save("docs/excel/CMAS-Score-Import-Template.xlsx")
print("Saved docs/excel/CMAS-Score-Import-Template.xlsx")
